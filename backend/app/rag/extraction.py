"""Text extraction from PDF, DOCX, TXT, and Markdown files.

Output is a list of TextBlocks that remember where the text came from (page number for
PDFs, heading path such as "Returns > Damaged items" for DOCX and Markdown). That
location is what citations show to the customer.
"""

import re
import zipfile
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path

import docx
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader
from pypdf.errors import PdfReadError

# A DOCX is a zip file; refuse ones that would inflate to an absurd size (zip bombs).
MAX_DOCX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024

_CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
_MD_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")
_MD_FENCE = re.compile(r"^\s*(```|~~~)")
_DOCX_HEADING_STYLE = re.compile(r"^Heading (\d)$")


class ExtractionError(Exception):
    """The file cannot be turned into text. The message is shown to admins."""


@dataclass
class TextBlock:
    text: str
    page_number: int | None = None
    section: str | None = None


@dataclass
class ExtractedDocument:
    blocks: list[TextBlock] = field(default_factory=list)
    page_count: int | None = None


def clean_text(text: str) -> str:
    """Normalise line endings, drop control characters, and squeeze blank lines."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _CONTROL_CHARS.sub("", text)
    lines = [line.rstrip() for line in text.split("\n")]
    text = "\n".join(lines)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def decode_text(data: bytes) -> str:
    """Decode a text file as UTF-8 (with or without BOM), falling back to Windows-1252."""
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ExtractionError("The file is not readable text (unknown character encoding).")


def _section_path(stack: list[tuple[int, str]]) -> str | None:
    return " > ".join(title for _, title in stack) or None


def _push_heading(stack: list[tuple[int, str]], level: int, title: str) -> None:
    while stack and stack[-1][0] >= level:
        stack.pop()
    stack.append((level, title))


# --------------------------------------------------------------------------- formats


def _extract_pdf(data: bytes) -> ExtractedDocument:
    try:
        reader = PdfReader(BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ExtractionError("The PDF is password-protected.")
        pages = list(reader.pages)
    except ExtractionError:
        raise
    except (PdfReadError, ValueError, KeyError, TypeError) as exc:
        raise ExtractionError("The PDF file is damaged or not a valid PDF.") from exc

    blocks = []
    for number, page in enumerate(pages, start=1):
        try:
            text = clean_text(page.extract_text() or "")
        except Exception as exc:  # noqa: BLE001 - pypdf raises many types on odd files
            raise ExtractionError(f"Could not read text on page {number}.") from exc
        if text:
            blocks.append(TextBlock(text=text, page_number=number))
    if not blocks:
        raise ExtractionError(
            "No text could be extracted. Scanned PDFs (images of text) are not supported."
        )
    return ExtractedDocument(blocks=blocks, page_count=len(pages))


def _docx_table_text(table: Table) -> str:
    rows = []
    for row in table.rows:
        cells = [" ".join(cell.text.split()) for cell in row.cells]
        if any(cells):
            rows.append(" | ".join(cells))
    return "\n".join(rows)


def _extract_docx(data: bytes) -> ExtractedDocument:
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            if sum(info.file_size for info in archive.infolist()) > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise ExtractionError("The DOCX file expands to an unreasonable size.")
        document = docx.Document(BytesIO(data))
    except ExtractionError:
        raise
    except Exception as exc:  # noqa: BLE001 - python-docx raises many types on bad input
        raise ExtractionError("The DOCX file is damaged or not a valid Word document.") from exc

    blocks: list[TextBlock] = []
    stack: list[tuple[int, str]] = []
    current: list[str] = []

    def flush() -> None:
        text = clean_text("\n".join(current))
        if text:
            blocks.append(TextBlock(text=text, section=_section_path(stack)))
        current.clear()

    for item in document.iter_inner_content():
        if isinstance(item, Paragraph):
            text = item.text.strip()
            style = item.style.name if item.style is not None else ""
            level_match = _DOCX_HEADING_STYLE.match(style)
            if text and (level_match or style == "Title"):
                flush()
                _push_heading(stack, int(level_match.group(1)) if level_match else 0, text)
                current.append(text)  # headings stay searchable
            elif text:
                current.append(text)
        elif isinstance(item, Table):
            table_text = _docx_table_text(item)
            if table_text:
                current.append(table_text)
    flush()
    if not blocks:
        raise ExtractionError("The Word document contains no text.")
    return ExtractedDocument(blocks=blocks)


def _extract_markdown(data: bytes) -> ExtractedDocument:
    blocks: list[TextBlock] = []
    stack: list[tuple[int, str]] = []
    current: list[str] = []
    in_code_block = False

    def flush() -> None:
        text = clean_text("\n".join(current))
        if text:
            blocks.append(TextBlock(text=text, section=_section_path(stack)))
        current.clear()

    for line in decode_text(data).splitlines():
        if _MD_FENCE.match(line):
            in_code_block = not in_code_block
            current.append(line)
            continue
        heading = None if in_code_block else _MD_HEADING.match(line)
        if heading:
            flush()
            title = heading.group(2).strip()
            _push_heading(stack, len(heading.group(1)), title)
            current.append(title)
        else:
            current.append(line)
    flush()
    if not blocks:
        raise ExtractionError("The Markdown file contains no text.")
    return ExtractedDocument(blocks=blocks)


def _extract_plain_text(data: bytes) -> ExtractedDocument:
    text = clean_text(decode_text(data))
    if not text:
        raise ExtractionError("The text file is empty.")
    return ExtractedDocument(blocks=[TextBlock(text=text)])


_EXTRACTORS = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".md": _extract_markdown,
    ".txt": _extract_plain_text,
}


def extract_bytes(data: bytes, extension: str) -> ExtractedDocument:
    try:
        extractor = _EXTRACTORS[extension.lower()]
    except KeyError:
        raise ExtractionError(f"Unsupported file type: {extension}") from None
    return extractor(data)


def extract_file(path: Path, extension: str) -> ExtractedDocument:
    try:
        data = path.read_bytes()
    except FileNotFoundError:
        raise ExtractionError("The stored file is missing; upload the document again.") from None
    return extract_bytes(data, extension)
