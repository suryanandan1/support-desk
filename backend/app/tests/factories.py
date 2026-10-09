"""Builders for real PDF and DOCX files used by the tests (no binary fixtures needed)."""

from io import BytesIO

import docx
from pypdf import PdfReader, PdfWriter


def _pdf_escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_pdf(pages: list[str]) -> bytes:
    """A minimal valid PDF with one text page per item (Helvetica, one line per \\n)."""
    objects: dict[int, bytes] = {
        1: b"<< /Type /Catalog /Pages 2 0 R >>",
        3: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    }
    page_ids = []
    next_id = 4
    for text in pages:
        page_id, content_id = next_id, next_id + 1
        next_id += 2
        operators = ["BT", "/F1 11 Tf", "14 TL", "72 720 Td"]
        operators += [f"({_pdf_escape(line)}) Tj T*" for line in text.split("\n")]
        operators.append("ET")
        stream = "\n".join(operators).encode("latin-1")
        objects[content_id] = b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream)
        objects[page_id] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content_id} 0 R >>"
        ).encode()
        page_ids.append(page_id)
    kids = " ".join(f"{pid} 0 R" for pid in page_ids)
    objects[2] = f"<< /Type /Pages /Kids [{kids}] /Count {len(page_ids)} >>".encode()

    out = BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = {}
    for object_id in sorted(objects):
        offsets[object_id] = out.tell()
        out.write(b"%d 0 obj\n%s\nendobj\n" % (object_id, objects[object_id]))
    xref_start = out.tell()
    count = max(objects) + 1
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % count)
    for object_id in range(1, count):
        out.write(b"%010d 00000 n \n" % offsets[object_id])
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (count, xref_start))
    return out.getvalue()


def make_encrypted_pdf(pages: list[str], password: str = "secret") -> bytes:
    writer = PdfWriter(clone_from=PdfReader(BytesIO(make_pdf(pages))))
    writer.encrypt(password)
    out = BytesIO()
    writer.write(out)
    return out.getvalue()


def make_docx(items: list[tuple[str, str]]) -> bytes:
    """Build a DOCX from (kind, text) pairs; kind is "h1", "h2", "p", or "table" (rows "a|b;c|d")."""
    document = docx.Document()
    for kind, text in items:
        if kind in ("h1", "h2", "h3"):
            document.add_heading(text, level=int(kind[1]))
        elif kind == "table":
            rows = [row.split("|") for row in text.split(";")]
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    table.cell(r, c).text = value
        else:
            document.add_paragraph(text)
    out = BytesIO()
    document.save(out)
    return out.getvalue()
